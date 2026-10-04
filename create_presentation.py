import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

WORKSPACE_DIR = r"c:\Users\Abdulaziz\Desktop\призентация"
ASSETS_DIR = os.path.join(WORKSPACE_DIR, "assets")

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    # Modern University Luxury Palette
    BG_DARK = RGBColor(10, 17, 40)        # #0A1128 Deep academic navy
    CARD_BG = RGBColor(17, 27, 59)        # #111B3B Elevated card fill
    CARD_BG_HOVER = RGBColor(23, 37, 78)  # Slightly lighter container
    CARD_BORDER = RGBColor(38, 55, 105)   # #263769 Subtle border
    
    ACCENT_BLUE = RGBColor(56, 189, 248)  # Sky blue
    ACCENT_CYAN = RGBColor(34, 211, 238)  # Bright cyan
    ACCENT_VIOLET = RGBColor(168, 85, 247)# Vibrant violet
    TEXT_EN = RGBColor(255, 255, 255)     # Crisp pure white
    TEXT_UZ = RGBColor(148, 163, 184)     # Slate 400
    TEXT_MUTED = RGBColor(100, 116, 139)  # Slate 500
    
    GREEN_ACCENT = RGBColor(52, 211, 153) # #34D399 Emerald
    GREEN_BG = RGBColor(12, 43, 36)
    ROSE_ACCENT = RGBColor(251, 113, 133) # #FB7185 Rose
    ROSE_BG = RGBColor(46, 18, 28)

    FONT_FAMILY = "Segoe UI"

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_DARK

    def add_header(slide, slide_num, en_title, uz_title, pill_text=None):
        # Top Meta Pill
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(8.0), Inches(0.35))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = pill_text or f"BILINGUAL PRESENTATION (EN / UZ) • SLIDE {slide_num:02d} OF 09"
        p_tag.font.name = FONT_FAMILY
        p_tag.font.size = Pt(10)
        p_tag.font.bold = True
        p_tag.font.color.rgb = ACCENT_BLUE

        # Main Title Box
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.733), Inches(0.95))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        # English Title
        p_en = tf.paragraphs[0]
        p_en.text = en_title
        p_en.font.name = FONT_FAMILY
        p_en.font.size = Pt(23)
        p_en.font.bold = True
        p_en.font.color.rgb = TEXT_EN
        p_en.space_after = Pt(2)
        
        # Uzbek Title
        p_uz = tf.add_paragraph()
        p_uz.text = uz_title
        p_uz.font.name = FONT_FAMILY
        p_uz.font.size = Pt(15)
        p_uz.font.italic = True
        p_uz.font.color.rgb = ACCENT_CYAN

        # Sleek Divider
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.75), Inches(11.733), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(30, 41, 75)
        line.line.color.rgb = RGBColor(30, 41, 75)

    # -------------------------------------------------------------
    # SLIDE 1: TITLE
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Big Card Container
    hero_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.75), Inches(11.733), Inches(6.0))
    hero_card.fill.solid()
    hero_card.fill.fore_color.rgb = CARD_BG
    hero_card.line.color.rgb = CARD_BORDER
    hero_card.line.width = Pt(1.5)

    # Left content area for slide 1
    t_box = slide1.shapes.add_textbox(Inches(1.4), Inches(1.4), Inches(6.6), Inches(4.8))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

    # Pill tag
    p_b = tf1.paragraphs[0]
    p_b.text = "UNIVERSITY PRESENTATION • BILINGUAL EDITION"
    p_b.font.name = FONT_FAMILY
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = ACCENT_BLUE
    p_b.space_after = Pt(16)

    # Title EN
    p1 = tf1.add_paragraph()
    p1.text = "Social Networks in My Life"
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_EN
    p1.space_after = Pt(4)

    # Title UZ
    p2 = tf1.add_paragraph()
    p2.text = "Ijtimoiy tarmoqlar mening hayotimda"
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(22)
    p2.font.italic = True
    p2.font.color.rgb = ACCENT_CYAN
    p2.space_after = Pt(24)

    # Subtitle EN
    p3 = tf1.add_paragraph()
    p3.text = "How social media affects communication, learning and everyday life"
    p3.font.name = FONT_FAMILY
    p3.font.size = Pt(16)
    p3.font.bold = True
    p3.font.color.rgb = RGBColor(226, 232, 240)
    p3.space_after = Pt(4)

    # Subtitle UZ
    p4 = tf1.add_paragraph()
    p4.text = "Ijtimoiy tarmoqlar muloqot, ta’lim va kundalik hayotimga qanday ta’sir qiladi"
    p4.font.name = FONT_FAMILY
    p4.font.size = Pt(13)
    p4.font.italic = True
    p4.font.color.rgb = TEXT_UZ
    p4.space_after = Pt(28)

    # Speaker footer info
    p5 = tf1.add_paragraph()
    p5.text = "English (Primary) + O‘zbekcha (Tarjima) • 9 Interactive Slides"
    p5.font.name = FONT_FAMILY
    p5.font.size = Pt(11)
    p5.font.color.rgb = TEXT_MUTED

    # Hero Network graphic on Right side
    hero_img_path = os.path.join(ASSETS_DIR, "hero_graphic.png")
    if os.path.exists(hero_img_path):
        slide1.shapes.add_picture(hero_img_path, Inches(8.1), Inches(1.25), width=Inches(4.1))

    # -------------------------------------------------------------
    # SLIDE 2: WHAT ARE SOCIAL NETWORKS?
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, 2, "What Are Social Networks?", "Ijtimoiy tarmoqlar nima?")

    # Big definition banner
    def_card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.95), Inches(11.733), Inches(1.8))
    def_card.fill.solid()
    def_card.fill.fore_color.rgb = CARD_BG
    def_card.line.color.rgb = CARD_BORDER
    def_card.line.width = Pt(1)

    tf_def = def_card.text_frame
    tf_def.word_wrap = True
    tf_def.margin_left = Inches(0.4)
    tf_def.margin_right = Inches(0.4)
    tf_def.margin_top = Inches(0.26)

    p_d1 = tf_def.paragraphs[0]
    p_d1.text = "Social networks are online platforms where people can communicate, share information and create content. They have become an important part of modern life."
    p_d1.font.name = FONT_FAMILY
    p_d1.font.size = Pt(17)
    p_d1.font.bold = True
    p_d1.font.color.rgb = TEXT_EN
    p_d1.space_after = Pt(8)

    p_d2 = tf_def.add_paragraph()
    p_d2.text = "Ijtimoiy tarmoqlar — odamlar muloqot qilishi, ma’lumot almashishi va kontent yaratishi mumkin bo‘lgan onlayn platformalardir. Ular zamonaviy hayotning muhim qismiga aylandi."
    p_d2.font.name = FONT_FAMILY
    p_d2.font.size = Pt(13)
    p_d2.font.italic = True
    p_d2.font.color.rgb = TEXT_UZ

    # 3 Points Columns
    points = [
        ("Communication", "Muloqot", "Connecting with people across borders and staying in touch with friends & family.", "Chegarasiz insonlar bilan bog‘lanish hamda do‘stlar va oila bilan doim aloqada bo‘lish.", ACCENT_BLUE, "icon_comm.png"),
        ("Information", "Ma’lumot", "Accessing daily news, scientific insights, and verified academic resources instantly.", "Kunlik yangiliklar, ilmiy manbalar va ishonchli bilimlarga bir zumda ega bo‘lish.", ACCENT_CYAN, "icon_info.png"),
        ("Content Sharing", "Kontent almashish", "Expressing ideas, sharing creative photography, video materials, and personal stories.", "O‘z g‘oyalarini ifodalash, ijodiy fotosuratlar, video materiallar va hikoyalar ulashish.", ACCENT_VIOLET, "icon_play.png")
    ]

    card_w = Inches(3.75)
    gap = Inches(0.24)
    start_x = Inches(0.8)

    for i, (en_p, uz_p, en_sub, uz_sub, color, icon_name) in enumerate(points):
        cx = start_x + i * (card_w + gap)
        c_shape = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(4.0), card_w, Inches(3.0))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = color
        c_shape.line.width = Pt(1.5)

        # Place icon
        icon_path = os.path.join(ASSETS_DIR, icon_name)
        if os.path.exists(icon_path):
            slide2.shapes.add_picture(icon_path, cx + Inches(0.3), Inches(4.25), width=Inches(0.65), height=Inches(0.65))

        tf_c = c_shape.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.3)
        tf_c.margin_right = Inches(0.3)
        tf_c.margin_top = Inches(1.05)

        p_pe = tf_c.paragraphs[0]
        p_pe.text = en_p
        p_pe.font.name = FONT_FAMILY
        p_pe.font.size = Pt(19)
        p_pe.font.bold = True
        p_pe.font.color.rgb = TEXT_EN
        p_pe.space_after = Pt(2)

        p_pu = tf_c.add_paragraph()
        p_pu.text = uz_p
        p_pu.font.name = FONT_FAMILY
        p_pu.font.size = Pt(13.5)
        p_pu.font.italic = True
        p_pu.font.color.rgb = color
        p_pu.space_after = Pt(10)

        p_se = tf_c.add_paragraph()
        p_se.text = en_sub
        p_se.font.name = FONT_FAMILY
        p_se.font.size = Pt(11.5)
        p_se.font.color.rgb = RGBColor(226, 232, 240)
        p_se.space_after = Pt(4)

        p_su = tf_c.add_paragraph()
        p_su.text = uz_sub
        p_su.font.name = FONT_FAMILY
        p_su.font.size = Pt(10)
        p_su.font.italic = True
        p_su.font.color.rgb = TEXT_UZ

    # -------------------------------------------------------------
    # SLIDE 3: SOCIAL NETWORKS I USE
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, 3, "Social Networks I Use", "Men foydalanadigan ijtimoiy tarmoqlar")

    # Introduction note
    note_box = slide3.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(11.733), Inches(0.4))
    tf_n = note_box.text_frame
    tf_n.word_wrap = True
    tf_n.margin_left = tf_n.margin_top = 0
    p_ne = tf_n.paragraphs[0]
    p_ne.text = "Examples of platforms commonly used by students / Talabalar hayotida keng foydalaniladigan platformalar misollari:"
    p_ne.font.name = FONT_FAMILY
    p_ne.font.size = Pt(12)
    p_ne.font.color.rgb = RGBColor(148, 163, 184)

    platforms = [
        ("Telegram",
         "communication and useful information",
         "muloqot va foydali ma’lumotlar uchun",
         "• Fast direct messaging and group discussions\n• University study channels and announcements\n• High-speed document and book sharing",
         "• Tezkor xabar almashish va guruh suhbatlari\n• Universitet o‘quv kanallari va e’lonlar\n• Kitoblar va hujjatlarni qulay almashish",
         RGBColor(0, 136, 204),
         "telegram.png"),
        ("Instagram",
         "photos, videos and creative content",
         "fotosuratlar, videolar va ijodiy kontent uchun",
         "• Visual storytelling and creative portfolio\n• Following educational designers and innovators\n• Short videos (Reels) showcasing modern ideas",
         "• Vizual hikoyalar va ijodiy ishlar namunasi\n• Ilhom beruvchi dizaynerlar va mutaxassislarni kuzatish\n• Zamonaviy g‘oyalarni aks ettiruvchi qisqa videolar",
         RGBColor(225, 48, 108),
         "instagram.png"),
        ("YouTube",
         "education and entertainment",
         "ta’lim va ko‘ngilochar kontent uchun",
         "• Comprehensive video lectures and tutorials\n• Science documentaries and academic channels\n• High-quality music and podcasts for relaxation",
         "• Chuqurlashtirilgan video darslar va ma’ruzalar\n• Ilmiy hujjatli filmlar va ta’limiy kanallar\n• Hordiq chiqarish uchun sifatli musiqa va podkastlar",
         RGBColor(255, 0, 0),
         "youtube.png")
    ]

    card_w = Inches(3.75)
    gap = Inches(0.24)
    start_x = Inches(0.8)

    for i, (name, en_desc, uz_desc, en_bullets, uz_bullets, brand_col, icon_file) in enumerate(platforms):
        cx = start_x + i * (card_w + gap)
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(2.35), card_w, Inches(4.7))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = brand_col
        card.line.width = Pt(1.5)

        # Place logo
        logo_path = os.path.join(ASSETS_DIR, icon_file)
        if os.path.exists(logo_path):
            slide3.shapes.add_picture(logo_path, cx + Inches(0.35), Inches(2.6), width=Inches(0.75), height=Inches(0.75))

        tf_c = card.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.35)
        tf_c.margin_right = Inches(0.35)
        tf_c.margin_top = Inches(1.15)

        # Platform Name
        p_name = tf_c.paragraphs[0]
        p_name.text = name
        p_name.font.name = FONT_FAMILY
        p_name.font.size = Pt(22)
        p_name.font.bold = True
        p_name.font.color.rgb = TEXT_EN
        p_name.space_after = Pt(4)

        # Subtitle EN
        p_se = tf_c.add_paragraph()
        p_se.text = en_desc
        p_se.font.name = FONT_FAMILY
        p_se.font.size = Pt(13)
        p_se.font.bold = True
        p_se.font.color.rgb = ACCENT_CYAN
        p_se.space_after = Pt(2)

        # Subtitle UZ
        p_su = tf_c.add_paragraph()
        p_su.text = uz_desc
        p_su.font.name = FONT_FAMILY
        p_su.font.size = Pt(11)
        p_su.font.italic = True
        p_su.font.color.rgb = TEXT_UZ
        p_su.space_after = Pt(12)

        # Bullets EN
        p_be = tf_c.add_paragraph()
        p_be.text = en_bullets
        p_be.font.name = FONT_FAMILY
        p_be.font.size = Pt(11)
        p_be.font.color.rgb = RGBColor(241, 245, 249)
        p_be.space_after = Pt(8)

        # Bullets UZ
        p_bu = tf_c.add_paragraph()
        p_bu.text = uz_bullets
        p_bu.font.name = FONT_FAMILY
        p_bu.font.size = Pt(9.5)
        p_bu.font.italic = True
        p_bu.font.color.rgb = RGBColor(148, 163, 184)

    # -------------------------------------------------------------
    # SLIDE 4: WHY DO I USE SOCIAL NETWORKS?
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, 4, "Why Do I Use Social Networks?", "Nega ijtimoiy tarmoqlardan foydalanaman?")

    # Statement Banner
    stmt_card = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.95), Inches(11.733), Inches(1.5))
    stmt_card.fill.solid()
    stmt_card.fill.fore_color.rgb = CARD_BG
    stmt_card.line.color.rgb = CARD_BORDER
    stmt_card.line.width = Pt(1)

    tf_s = stmt_card.text_frame
    tf_s.word_wrap = True
    tf_s.margin_left = Inches(0.4)
    tf_s.margin_right = Inches(0.4)
    tf_s.margin_top = Inches(0.22)

    p_se = tf_s.paragraphs[0]
    p_se.text = "I use social networks for several reasons. They help me communicate with other people, find useful information, learn new things and relax in my free time."
    p_se.font.name = FONT_FAMILY
    p_se.font.size = Pt(16.5)
    p_se.font.bold = True
    p_se.font.color.rgb = TEXT_EN
    p_se.space_after = Pt(6)

    p_su = tf_s.add_paragraph()
    p_su.text = "Men ijtimoiy tarmoqlardan bir nechta sabablar tufayli foydalanaman. Ular menga boshqa odamlar bilan muloqot qilish, foydali ma’lumot topish, yangi narsalarni o‘rganish va bo‘sh vaqtimda dam olishga yordam beradi."
    p_su.font.name = FONT_FAMILY
    p_su.font.size = Pt(12.5)
    p_su.font.italic = True
    p_su.font.color.rgb = TEXT_UZ

    # 4 Visual categories
    categories = [
        ("Communication", "Muloqot", "Staying in touch with friends, family, and classmates seamlessly.", "Do‘stlar, oila va kursdoshlar bilan oson va yaqin aloqada bo‘lish.", ACCENT_BLUE, "icon_comm.png"),
        ("Information", "Ma’lumot", "Discovering news, campus updates, and global affairs.", "Yangiliklar, universitet xabarlari va jahon voqealarini bilib borish.", ACCENT_CYAN, "icon_info.png"),
        ("Learning", "O‘rganish", "Acquiring practical skills, foreign languages, and study advice.", "Amaliy ko‘nikmalar, chet tillari va o‘quv maslahatlarini o‘zlashtirish.", ACCENT_VIOLET, "icon_learn.png"),
        ("Entertainment", "Ko‘ngilochar", "Listening to favorite music, watching videos, and resting peacefully.", "Sevimli musiqa tinglash, videolar ko‘rish va maroqli dam olish.", GREEN_ACCENT, "icon_play.png")
    ]

    cat_w = Inches(2.78)
    cat_gap = Inches(0.2)
    start_cx = Inches(0.8)

    for i, (en_c, uz_c, en_d, uz_d, color, icon_file) in enumerate(categories):
        x = start_cx + i * (cat_w + cat_gap)
        c_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(3.7), cat_w, Inches(3.3))
        c_box.fill.solid()
        c_box.fill.fore_color.rgb = CARD_BG
        c_box.line.color.rgb = color
        c_box.line.width = Pt(1.5)

        # Place mini icon
        ic_p = os.path.join(ASSETS_DIR, icon_file)
        if os.path.exists(ic_p):
            slide4.shapes.add_picture(ic_p, x + Inches(0.25), Inches(3.9), width=Inches(0.55), height=Inches(0.55))

        tf_cb = c_box.text_frame
        tf_cb.word_wrap = True
        tf_cb.margin_left = Inches(0.25)
        tf_cb.margin_right = Inches(0.25)
        tf_cb.margin_top = Inches(0.85)

        p_ce = tf_cb.paragraphs[0]
        p_ce.text = en_c
        p_ce.font.name = FONT_FAMILY
        p_ce.font.size = Pt(17)
        p_ce.font.bold = True
        p_ce.font.color.rgb = TEXT_EN
        p_ce.space_after = Pt(2)

        p_cu = tf_cb.add_paragraph()
        p_cu.text = uz_c
        p_cu.font.name = FONT_FAMILY
        p_cu.font.size = Pt(13)
        p_cu.font.italic = True
        p_cu.font.color.rgb = color
        p_cu.space_after = Pt(10)

        p_de = tf_cb.add_paragraph()
        p_de.text = en_d
        p_de.font.name = FONT_FAMILY
        p_de.font.size = Pt(11)
        p_de.font.color.rgb = RGBColor(226, 232, 240)
        p_de.space_after = Pt(4)

        p_du = tf_cb.add_paragraph()
        p_du.text = uz_d
        p_du.font.name = FONT_FAMILY
        p_du.font.size = Pt(9.5)
        p_du.font.italic = True
        p_du.font.color.rgb = TEXT_UZ

    # -------------------------------------------------------------
    # SLIDE 5: COMMUNICATION WITH PEOPLE
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, 5, "Communication with People", "Odamlar bilan muloqot")

    # Left Container (Core message)
    left_card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.95), Inches(5.7), Inches(5.1))
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = CARD_BG
    left_card.line.color.rgb = ACCENT_BLUE
    left_card.line.width = Pt(1.5)

    tf_lc = left_card.text_frame
    tf_lc.word_wrap = True
    tf_lc.margin_left = Inches(0.4)
    tf_lc.margin_right = Inches(0.4)
    tf_lc.margin_top = Inches(0.35)

    p_lct = tf_lc.paragraphs[0]
    p_lct.text = "CORE ADVANTAGE • ASOSIY AFZALLIK"
    p_lct.font.name = FONT_FAMILY
    p_lct.font.size = Pt(11)
    p_lct.font.bold = True
    p_lct.font.color.rgb = ACCENT_BLUE
    p_lct.space_after = Pt(14)

    p_lce = tf_lc.add_paragraph()
    p_lce.text = "Social networks make communication faster and easier. I can send messages, share photos and talk to friends and family even when they are far away."
    p_lce.font.name = FONT_FAMILY
    p_lce.font.size = Pt(17)
    p_lce.font.bold = True
    p_lce.font.color.rgb = TEXT_EN
    p_lce.space_after = Pt(12)

    p_lcu = tf_lc.add_paragraph()
    p_lcu.text = "Ijtimoiy tarmoqlar muloqotni tezroq va osonroq qiladi. Men do‘stlarim va oilam uzoqda bo‘lsa ham, xabar yuborishim, fotosurat ulashishim va ular bilan suhbatlashishim mumkin."
    p_lcu.font.name = FONT_FAMILY
    p_lcu.font.size = Pt(13)
    p_lcu.font.italic = True
    p_lcu.font.color.rgb = TEXT_UZ
    p_lcu.space_after = Pt(20)

    # Mini visual on bottom left
    comm_vis_path = os.path.join(ASSETS_DIR, "comm_visual.png")
    if os.path.exists(comm_vis_path):
        slide5.shapes.add_picture(comm_vis_path, Inches(1.5), Inches(4.55), width=Inches(3.8))

    # Right side 3 Pillars
    comm_aspects = [
        ("Instant Messaging & Voice Notes", "Tezkor xabarlar va ovozli suhbatlar",
         "Messages are delivered instantly, making everyday coordination effortless.",
         "Xabarlar bir lahzada yetib boradi va kundalik masalalarni hal qilishni osonlashtiradi.", ACCENT_BLUE),
        ("Sharing Photos & Memories", "Fotosuratlar va xotiralarni ulashish",
         "Sharing life updates and experiences with loved ones regardless of physical distance.",
         "Masofadan qat’i nazar, yaqin insonlar bilan hayotdagi qiziqarli voqealarni bo‘lishish.", ACCENT_CYAN),
        ("Global Video Calls", "Global video muloqot",
         "Seeing faces in real time bridges gaps across continents and time zones.",
         "Yuzma-yuz gaplashish orqali qit’alar va vaqt zonalarini bir zumda bog‘lash.", ACCENT_VIOLET)
    ]

    r_start_y = Inches(1.95)
    r_h = Inches(1.55)
    r_gap = Inches(0.22)

    for i, (en_a, uz_a, en_sub, uz_sub, col) in enumerate(comm_aspects):
        y = r_start_y + i * (r_h + r_gap)
        rc = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), y, Inches(5.733), r_h)
        rc.fill.solid()
        rc.fill.fore_color.rgb = CARD_BG
        rc.line.color.rgb = CARD_BORDER
        rc.line.width = Pt(1)

        tf_rc = rc.text_frame
        tf_rc.word_wrap = True
        tf_rc.margin_left = Inches(0.35)
        tf_rc.margin_right = Inches(0.35)
        tf_rc.margin_top = Inches(0.2)

        p_e = tf_rc.paragraphs[0]
        p_e.text = en_a
        p_e.font.name = FONT_FAMILY
        p_e.font.size = Pt(15)
        p_e.font.bold = True
        p_e.font.color.rgb = TEXT_EN
        p_e.space_after = Pt(2)

        p_u = tf_rc.add_paragraph()
        p_u.text = uz_a
        p_u.font.name = FONT_FAMILY
        p_u.font.size = Pt(12)
        p_u.font.italic = True
        p_u.font.color.rgb = col
        p_u.space_after = Pt(6)

        p_sub = tf_rc.add_paragraph()
        p_sub.text = en_sub
        p_sub.font.name = FONT_FAMILY
        p_sub.font.size = Pt(10.5)
        p_sub.font.color.rgb = RGBColor(226, 232, 240)
        p_sub.space_after = Pt(2)

        p_subu = tf_rc.add_paragraph()
        p_subu.text = uz_sub
        p_subu.font.name = FONT_FAMILY
        p_subu.font.size = Pt(9.5)
        p_subu.font.italic = True
        p_subu.font.color.rgb = TEXT_UZ

    # -------------------------------------------------------------
    # SLIDE 6: LEARNING AND EDUCATION
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, 6, "Learning and Education", "O‘rganish va ta’lim")

    # Core message banner
    ed_banner = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.95), Inches(11.733), Inches(1.5))
    ed_banner.fill.solid()
    ed_banner.fill.fore_color.rgb = CARD_BG
    ed_banner.line.color.rgb = CARD_BORDER
    ed_banner.line.width = Pt(1)

    tf_ed = ed_banner.text_frame
    tf_ed.word_wrap = True
    tf_ed.margin_left = Inches(0.4)
    tf_ed.margin_right = Inches(0.4)
    tf_ed.margin_top = Inches(0.22)

    p_ede = tf_ed.paragraphs[0]
    p_ede.text = "Social networks can also be useful for education. I can watch educational videos, follow useful channels and find new information about different topics."
    p_ede.font.name = FONT_FAMILY
    p_ede.font.size = Pt(16.5)
    p_ede.font.bold = True
    p_ede.font.color.rgb = TEXT_EN
    p_ede.space_after = Pt(6)

    p_edu = tf_ed.add_paragraph()
    p_edu.text = "Ijtimoiy tarmoqlar ta’lim uchun ham foydali bo‘lishi mumkin. Men ta’limiy videolar ko‘rishim, foydali kanallarni kuzatishim va turli mavzular haqida yangi ma’lumot topishim mumkin."
    p_edu.font.name = FONT_FAMILY
    p_edu.font.size = Pt(12.5)
    p_edu.font.italic = True
    p_edu.font.color.rgb = TEXT_UZ

    # 4 Visual educational elements
    edu_items = [
        ("Video lessons", "Video darslar",
         "Visual explanations and step-by-step guides for complex academic topics.",
         "Murakkab fanlar bo‘yicha ko‘rgazmali va tushunarli video darslar.", ACCENT_BLUE),
        ("Online courses", "Onlayn kurslar",
         "Structured modules to study languages, coding, design, and science.",
         "Chet tillari, dasturlash va dizaynni o‘rganish uchun tizimli kurslar.", ACCENT_CYAN),
        ("Educational channels", "Ta’limiy kanallar",
         "Expert curated posts, daily tips, and knowledge tests on social channels.",
         "Mutaxassislar tomonidan yuritiladigan tahliliy postlar va testlar.", ACCENT_VIOLET),
        ("Useful information", "Foydali ma’lumotlar",
         "Direct access to study guides, scientific articles, and learning groups.",
         "Ilmiy maqolalar, o‘quv qo‘llanmalari va foydali hamjamiyatlarga tezkor kirish.", GREEN_ACCENT)
    ]

    card_w = Inches(2.78)
    cat_gap = Inches(0.2)
    start_cx = Inches(0.8)

    for i, (en_t, uz_t, en_d, uz_d, col) in enumerate(edu_items):
        x = start_cx + i * (card_w + cat_gap)
        eb = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(3.7), cat_w, Inches(3.3))
        eb.fill.solid()
        eb.fill.fore_color.rgb = CARD_BG
        eb.line.color.rgb = col
        eb.line.width = Pt(1.5)

        tf_eb = eb.text_frame
        tf_eb.word_wrap = True
        tf_eb.margin_left = Inches(0.25)
        tf_eb.margin_right = Inches(0.25)
        tf_eb.margin_top = Inches(0.3)

        p_t1 = tf_eb.paragraphs[0]
        p_t1.text = f"PILLAR {i+1:02d}"
        p_t1.font.name = FONT_FAMILY
        p_t1.font.size = Pt(11)
        p_t1.font.bold = True
        p_t1.font.color.rgb = col
        p_t1.space_after = Pt(6)

        p_t2 = tf_eb.add_paragraph()
        p_t2.text = en_t
        p_t2.font.name = FONT_FAMILY
        p_t2.font.size = Pt(17)
        p_t2.font.bold = True
        p_t2.font.color.rgb = TEXT_EN
        p_t2.space_after = Pt(2)

        p_t3 = tf_eb.add_paragraph()
        p_t3.text = uz_t
        p_t3.font.name = FONT_FAMILY
        p_t3.font.size = Pt(13)
        p_t3.font.italic = True
        p_t3.font.color.rgb = col
        p_t3.space_after = Pt(10)

        p_t4 = tf_eb.add_paragraph()
        p_t4.text = en_d
        p_t4.font.name = FONT_FAMILY
        p_t4.font.size = Pt(11)
        p_t4.font.color.rgb = RGBColor(226, 232, 240)
        p_t4.space_after = Pt(4)

        p_t5 = tf_eb.add_paragraph()
        p_t5.text = uz_d
        p_t5.font.name = FONT_FAMILY
        p_t5.font.size = Pt(9.5)
        p_t5.font.italic = True
        p_t5.font.color.rgb = TEXT_UZ

    # -------------------------------------------------------------
    # SLIDE 7: ENTERTAINMENT AND FREE TIME
    # -------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, 7, "Entertainment and Free Time", "Ko‘ngilochar va bo‘sh vaqt")

    # Main text card
    ent_card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.95), Inches(11.733), Inches(1.7))
    ent_card.fill.solid()
    ent_card.fill.fore_color.rgb = CARD_BG
    ent_card.line.color.rgb = CARD_BORDER
    ent_card.line.width = Pt(1)

    tf_ent = ent_card.text_frame
    tf_ent.word_wrap = True
    tf_ent.margin_left = Inches(0.4)
    tf_ent.margin_right = Inches(0.4)
    tf_ent.margin_top = Inches(0.24)

    p_ee = tf_ent.paragraphs[0]
    p_ee.text = "Social networks are also a way to relax. I can watch videos, listen to music, see interesting content and discover new ideas. However, I try not to spend too much time online."
    p_ee.font.name = FONT_FAMILY
    p_ee.font.size = Pt(16.5)
    p_ee.font.bold = True
    p_ee.font.color.rgb = TEXT_EN
    p_ee.space_after = Pt(6)

    p_eu = tf_ent.add_paragraph()
    p_eu.text = "Ijtimoiy tarmoqlar dam olish usuli hamdir. Men videolar ko‘rishim, musiqa tinglashim, qiziqarli kontent tomosha qilishim va yangi g‘oyalarni kashf qilishim mumkin. Biroq internetda juda ko‘p vaqt sarflamaslikka harakat qilaman."
    p_eu.font.name = FONT_FAMILY
    p_eu.font.size = Pt(13)
    p_eu.font.italic = True
    p_eu.font.color.rgb = TEXT_UZ

    # 3 Pillars below
    pillars = [
        ("Relaxation & Music", "Dam olish va musiqa",
         "Unwinding after university lectures with calming music and inspirational videos.",
         "Ma’ruzalardan so‘ng yoqimli musiqa va qiziqarli videolar bilan charchoqni chiqarish.", ACCENT_BLUE),
        ("Discovering Creative Ideas", "Yangi ijodiy g‘oyalar",
         "Finding creative DIY projects, photography tips, and innovative concepts.",
         "Ijodiy loyihalar, yangi fotog‘oyalar va qiziqarli mashg‘ulotlarni kashf etish.", ACCENT_VIOLET),
        ("Conscious Screen Time", "Ongli me’yor va nazorat",
         "Preventing excessive passive scrolling to keep time for offline life and health.",
         "Haqiqiy hayot va salomatlik uchun virtual olamda sarflanadigan vaqtni me’yorida ushlash.", GREEN_ACCENT)
    ]

    card_w = Inches(3.75)
    gap = Inches(0.24)
    start_x = Inches(0.8)

    for i, (en_p, uz_p, en_d, uz_d, col) in enumerate(pillars):
        cx = start_x + i * (card_w + gap)
        c_shape = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(3.9), card_w, Inches(3.1))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = col
        c_shape.line.width = Pt(1.5)

        tf_c = c_shape.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.3)
        tf_c.margin_right = Inches(0.3)
        tf_c.margin_top = Inches(0.3)

        p_num = tf_c.paragraphs[0]
        p_num.text = f"ASPECT 0{i+1}"
        p_num.font.name = FONT_FAMILY
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = col
        p_num.space_after = Pt(4)

        p_pe = tf_c.add_paragraph()
        p_pe.text = en_p
        p_pe.font.name = FONT_FAMILY
        p_pe.font.size = Pt(18)
        p_pe.font.bold = True
        p_pe.font.color.rgb = TEXT_EN
        p_pe.space_after = Pt(2)

        p_pu = tf_c.add_paragraph()
        p_pu.text = uz_p
        p_pu.font.name = FONT_FAMILY
        p_pu.font.size = Pt(13)
        p_pu.font.italic = True
        p_pu.font.color.rgb = col
        p_pu.space_after = Pt(10)

        p_se = tf_c.add_paragraph()
        p_se.text = en_d
        p_se.font.name = FONT_FAMILY
        p_se.font.size = Pt(11.5)
        p_se.font.color.rgb = RGBColor(226, 232, 240)
        p_se.space_after = Pt(4)

        p_su = tf_c.add_paragraph()
        p_su.text = uz_d
        p_su.font.name = FONT_FAMILY
        p_su.font.size = Pt(10)
        p_su.font.italic = True
        p_su.font.color.rgb = TEXT_UZ

    # -------------------------------------------------------------
    # SLIDE 8: ADVANTAGES AND DISADVANTAGES
    # -------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, 8, "Advantages and Disadvantages", "Afzalliklari va kamchiliklari")

    col_w = Inches(5.72)
    
    # Left Column: ADVANTAGES
    adv_box = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.95), col_w, Inches(5.1))
    adv_box.fill.solid()
    adv_box.fill.fore_color.rgb = CARD_BG
    adv_box.line.color.rgb = GREEN_ACCENT
    adv_box.line.width = Pt(1.5)

    # Check badge
    chk_p = os.path.join(ASSETS_DIR, "icon_check.png")
    if os.path.exists(chk_p):
        slide8.shapes.add_picture(chk_p, Inches(1.15), Inches(2.2), width=Inches(0.45), height=Inches(0.45))

    tf_a = adv_box.text_frame
    tf_a.word_wrap = True
    tf_a.margin_left = Inches(0.4)
    tf_a.margin_right = Inches(0.4)
    tf_a.margin_top = Inches(0.3)

    p_ah = tf_a.paragraphs[0]
    p_ah.text = "       ADVANTAGES"
    p_ah.font.name = FONT_FAMILY
    p_ah.font.size = Pt(19)
    p_ah.font.bold = True
    p_ah.font.color.rgb = GREEN_ACCENT
    p_ah.space_after = Pt(1)

    p_ahu = tf_a.add_paragraph()
    p_ahu.text = "       AFZALLIKLARI"
    p_ahu.font.name = FONT_FAMILY
    p_ahu.font.size = Pt(13)
    p_ahu.font.italic = True
    p_ahu.font.color.rgb = RGBColor(167, 243, 208)
    p_ahu.space_after = Pt(18)

    adv_items = [
        ("Easy communication", "Oson muloqot"),
        ("Quick access to information", "Ma’lumotga tezkor kirish"),
        ("Educational opportunities", "Ta’lim imkoniyatlari"),
        ("Entertainment", "Ko‘ngilochar imkoniyatlar")
    ]

    for en_item, uz_item in adv_items:
        p_i = tf_a.add_paragraph()
        p_i.text = f"✔  {en_item}"
        p_i.font.name = FONT_FAMILY
        p_i.font.size = Pt(15.5)
        p_i.font.bold = True
        p_i.font.color.rgb = TEXT_EN
        p_i.space_after = Pt(1)

        p_iu = tf_a.add_paragraph()
        p_iu.text = f"     {uz_item}"
        p_iu.font.name = FONT_FAMILY
        p_iu.font.size = Pt(12.5)
        p_iu.font.italic = True
        p_iu.font.color.rgb = TEXT_UZ
        p_iu.space_after = Pt(11)

    # Right Column: DISADVANTAGES
    dis_box = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.95), col_w, Inches(5.1))
    dis_box.fill.solid()
    dis_box.fill.fore_color.rgb = CARD_BG
    dis_box.line.color.rgb = ROSE_ACCENT
    dis_box.line.width = Pt(1.5)

    # Cross badge
    crs_p = os.path.join(ASSETS_DIR, "icon_cross.png")
    if os.path.exists(crs_p):
        slide8.shapes.add_picture(crs_p, Inches(7.15), Inches(2.2), width=Inches(0.45), height=Inches(0.45))

    tf_d = dis_box.text_frame
    tf_d.word_wrap = True
    tf_d.margin_left = Inches(0.4)
    tf_d.margin_right = Inches(0.4)
    tf_d.margin_top = Inches(0.3)

    p_dh = tf_d.paragraphs[0]
    p_dh.text = "       DISADVANTAGES"
    p_dh.font.name = FONT_FAMILY
    p_dh.font.size = Pt(19)
    p_dh.font.bold = True
    p_dh.font.color.rgb = ROSE_ACCENT
    p_dh.space_after = Pt(1)

    p_dhu = tf_d.add_paragraph()
    p_dhu.text = "       KAMCHILIKLARI"
    p_dhu.font.name = FONT_FAMILY
    p_dhu.font.size = Pt(13)
    p_dhu.font.italic = True
    p_dhu.font.color.rgb = RGBColor(254, 205, 211)
    p_dhu.space_after = Pt(18)

    dis_items = [
        ("Too much screen time", "Ekran qarshisida ortiqcha vaqt"),
        ("Distraction", "Diqqatni chalg‘itishi"),
        ("False information", "Noto‘g‘ri ma’lumotlar"),
        ("Privacy problems", "Maxfiylik muammolari")
    ]

    for en_item, uz_item in dis_items:
        p_i = tf_d.add_paragraph()
        p_i.text = f"✖  {en_item}"
        p_i.font.name = FONT_FAMILY
        p_i.font.size = Pt(15.5)
        p_i.font.bold = True
        p_i.font.color.rgb = TEXT_EN
        p_i.space_after = Pt(1)

        p_iu = tf_d.add_paragraph()
        p_iu.text = f"     {uz_item}"
        p_iu.font.name = FONT_FAMILY
        p_iu.font.size = Pt(12.5)
        p_iu.font.italic = True
        p_iu.font.color.rgb = TEXT_UZ
        p_iu.space_after = Pt(11)

    # -------------------------------------------------------------
    # SLIDE 9: CONCLUSION
    # -------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, 9, "Conclusion", "Xulosa")

    # Main synthesis card
    conc_card = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.95), Inches(11.733), Inches(2.25))
    conc_card.fill.solid()
    conc_card.fill.fore_color.rgb = CARD_BG
    conc_card.line.color.rgb = ACCENT_BLUE
    conc_card.line.width = Pt(1.5)

    tf_c9 = conc_card.text_frame
    tf_c9.word_wrap = True
    tf_c9.margin_left = Inches(0.45)
    tf_c9.margin_right = Inches(0.45)
    tf_c9.margin_top = Inches(0.28)

    p_c9e = tf_c9.paragraphs[0]
    p_c9e.text = "Social networks are an important part of modern life. They help me communicate, learn and relax. However, I believe it is important to use them responsibly and keep a healthy balance between online and real life."
    p_c9e.font.name = FONT_FAMILY
    p_c9e.font.size = Pt(17)
    p_c9e.font.bold = True
    p_c9e.font.color.rgb = TEXT_EN
    p_c9e.space_after = Pt(8)

    p_c9u = tf_c9.add_paragraph()
    p_c9u.text = "Ijtimoiy tarmoqlar zamonaviy hayotning muhim qismidir. Ular menga muloqot qilish, o‘rganish va dam olishga yordam beradi. Biroq ulardan mas’uliyat bilan foydalanish va onlayn hamda real hayot o‘rtasida sog‘lom muvozanatni saqlash muhim deb hisoblayman."
    p_c9u.font.name = FONT_FAMILY
    p_c9u.font.size = Pt(13)
    p_c9u.font.italic = True
    p_c9u.font.color.rgb = TEXT_UZ

    # 3 Takeaway summary pills
    takeaways = [
        ("Responsible Usage", "Mas’uliyatli foydalanish", ACCENT_BLUE),
        ("Focused Purpose", "Aniq maqsad bilan yondashuv", ACCENT_CYAN),
        ("Healthy Life Balance", "Sog‘lom hayot muvozanati", ACCENT_VIOLET)
    ]
    tw_w = Inches(3.75)
    tw_gap = Inches(0.24)
    for i, (en_tw, uz_tw, col) in enumerate(takeaways):
        tx = start_x + i * (tw_w + tw_gap)
        tw_box = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, tx, Inches(4.4), tw_w, Inches(0.85))
        tw_box.fill.solid()
        tw_box.fill.fore_color.rgb = CARD_BG
        tw_box.line.color.rgb = col
        tw_box.line.width = Pt(1)

        tf_tw = tw_box.text_frame
        tf_tw.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf_tw.margin_left = tf_tw.margin_right = Inches(0.2)
        p_te = tf_tw.paragraphs[0]
        p_te.text = f"★  {en_tw}"
        p_te.font.name = FONT_FAMILY
        p_te.font.size = Pt(14)
        p_te.font.bold = True
        p_te.font.color.rgb = TEXT_EN
        p_te.alignment = PP_ALIGN.CENTER
        
        p_tu = tf_tw.add_paragraph()
        p_tu.text = uz_tw
        p_tu.font.name = FONT_FAMILY
        p_tu.font.size = Pt(11)
        p_tu.font.italic = True
        p_tu.font.color.rgb = col
        p_tu.alignment = PP_ALIGN.CENTER

    # Final Thank You Card
    ty_card = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.45), Inches(11.733), Inches(1.55))
    ty_card.fill.solid()
    ty_card.fill.fore_color.rgb = RGBColor(16, 37, 77)
    ty_card.line.color.rgb = ACCENT_CYAN
    ty_card.line.width = Pt(1.5)

    tf_ty = ty_card.text_frame
    tf_ty.word_wrap = True
    tf_ty.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_ty.margin_left = Inches(0.4)
    tf_ty.margin_right = Inches(0.4)

    p_tye = tf_ty.paragraphs[0]
    p_tye.text = "Thank you for your attention!"
    p_tye.font.name = FONT_FAMILY
    p_tye.font.size = Pt(24)
    p_tye.font.bold = True
    p_tye.font.color.rgb = TEXT_EN
    p_tye.alignment = PP_ALIGN.CENTER
    p_tye.space_after = Pt(3)

    p_tyu = tf_ty.add_paragraph()
    p_tyu.text = "E’tiboringiz uchun rahmat!"
    p_tyu.font.name = FONT_FAMILY
    p_tyu.font.size = Pt(17)
    p_tyu.font.italic = True
    p_tyu.font.color.rgb = ACCENT_CYAN
    p_tyu.alignment = PP_ALIGN.CENTER

    output_path = os.path.join(WORKSPACE_DIR, "Social_Networks_in_My_Life.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    build_presentation()
