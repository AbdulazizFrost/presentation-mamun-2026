from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[6])
sld = slide._element

# Transition XML for smooth fade
trans_xml = parse_xml(f'<p:transition {nsdecls("p")} spd="med" advClick="1"><p:fade/></p:transition>')
sld.append(trans_xml)

test_path = r'C:\Users\Abdulaziz\.gemini\antigravity\brain\acab1f40-302f-44eb-b996-e481750ad49f\test_trans.pptx'
prs.save(test_path)
print("Successfully added transition!")
