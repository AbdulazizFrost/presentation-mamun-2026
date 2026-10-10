"""
Build Time_Management_for_Students.pptx from the web slides.

    python build_pptx.py

Each slide is printed by headless Microsoft Edge (the same 1200x675 page the
"Export PDF" button produces), rendered to a PNG and placed full-screen on a
16:9 PowerPoint slide. The speech for every slide (title EN, title UZ, Uzbek
talk) is copied from index.html into the PowerPoint speaker notes.

Needs: python-pptx, pymupdf (pip install python-pptx pymupdf) and Microsoft Edge.
"""
import functools
import json
import os
import re
import subprocess
import tempfile
import threading
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer

import pymupdf
from pptx import Presentation
from pptx.util import Emu

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, "Time_Management_for_Students.pptx")
EDGE_PATHS = [
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
]


def read_notes():
    html = open(os.path.join(ROOT, "index.html"), encoding="utf-8").read()
    block = html[html.index("const speakerNotes = {"):]
    block = block[:block.index("    };")]
    fields = re.findall(r'(titleEn|titleUz|uz): (".*?")(?:,|\n)', block)
    values = [json.loads(v) for _, v in fields]
    return [values[i:i + 3] for i in range(0, len(values), 3)]


def print_pdf(pdf_path):
    edge = next((p for p in EDGE_PATHS if os.path.exists(p)), None)
    if not edge:
        raise SystemExit("Microsoft Edge not found")
    handler = functools.partial(SimpleHTTPRequestHandler, directory=ROOT)
    server = ThreadingHTTPServer(("127.0.0.1", 0), handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        with tempfile.TemporaryDirectory() as profile:
            subprocess.run([
                edge, "--headless=new", "--disable-gpu", f"--user-data-dir={profile}",
                "--no-pdf-header-footer", "--virtual-time-budget=5000",
                f"--print-to-pdf={pdf_path}", f"http://127.0.0.1:{server.server_port}/index.html?print",
            ], check=True, timeout=120)
    finally:
        server.shutdown()


def main():
    notes = read_notes()
    with tempfile.TemporaryDirectory() as tmp:
        pdf_path = os.path.join(tmp, "deck.pdf")
        print_pdf(pdf_path)
        doc = pymupdf.open(pdf_path)

        prs = Presentation()
        prs.slide_width = Emu(12192000)   # 13.333 in — 16:9
        prs.slide_height = Emu(6858000)   # 7.5 in
        blank = prs.slide_layouts[6]
        for i, page in enumerate(doc):
            png = os.path.join(tmp, f"slide{i + 1}.png")
            page.get_pixmap(dpi=192).save(png)
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(png, 0, 0, prs.slide_width, prs.slide_height)
            if i < len(notes):
                title_en, title_uz, talk = notes[i]
                slide.notes_slide.notes_text_frame.text = (
                    f"1️⃣ Sarlavha (avval inglizcha, keyin o‘zbekcha):\n🇬🇧 {title_en}\n🇺🇿 {title_uz}"
                    f"\n\n2️⃣ Nutq (o‘zbekcha):\n{talk}"
                )
        doc.close()
        prs.save(OUT)
    print(f"Saved {OUT} ({len(notes)} slides with notes)")


if __name__ == "__main__":
    main()
