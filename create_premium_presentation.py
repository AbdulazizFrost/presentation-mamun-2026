import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

WORKSPACE_DIR = r"c:\Users\Abdulaziz\Desktop\призентация"
ASSETS_DIR = os.path.join(WORKSPACE_DIR, "assets_premium")

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    FONT_FAMILY = "Segoe UI"

    # Unified Cohesive Dark Navy Luxury Theme (Across ALL 9 Slides)
    BG_DARK = RGBColor(9, 14, 26)          # #090E1A Deep midnight navy
    CARD_BG = RGBColor(17, 24, 39)         # #111827 Elevated dark slate card
    CARD_BORDER = RGBColor(30, 41, 59)     # #1E293B Clean card outline
    DIVIDER = RGBColor(30, 41, 65)         # Subtle rule line

    TEXT_EN = RGBColor(255, 255, 255)      # Primary crisp white for English
    TEXT_UZ = RGBColor(148, 163, 184)      # Secondary soft slate-400 for Uzbek
    ACCENT_BLUE = RGBColor(56, 189, 248)   # Sky-400 electric blue
    ACCENT_CYAN = RGBColor(34, 211, 238)   # Cyan highlight
    ACCENT_GREEN = RGBColor(52, 211, 153)  # Emerald for advantages
    ACCENT_ROSE = RGBColor(251, 113, 133)  # Rose for disadvantages

    def set_slide_bg(slide):
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = BG_DARK

        # Add smooth fade transition for PowerPoint presentation mode
        from pptx.oxml import parse_xml
        from pptx.oxml.ns import nsdecls
        trans_xml = parse_xml(f'<p:transition {nsdecls("p")} spd="med" advClick="1"><p:fade/></p:transition>')
        slide._element.append(trans_xml)

    def add_slide_header(slide, slide_num, en_title, uz_title, tag_text=None):
        # Category Tag
        tag_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.42), Inches(8.0), Inches(0.3))
        tf_t = tag_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = tag_text or f"UNIVERSITY PRESENTATION (EN / UZ) • SLIDE {slide_num:02d} OF 09"
        p_t.font.name = FONT_FAMILY
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = ACCENT_BLUE

        # Title Box
        title_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.72), Inches(11.533), Inches(0.88))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_en = tf.paragraphs[0]
        p_en.text = en_title
        p_en.font.name = FONT_FAMILY
        p_en.font.size = Pt(24)
        p_en.font.bold = True
        p_en.font.color.rgb = TEXT_EN
        p_en.space_after = Pt(2)

        p_uz = tf.add_paragraph()
        p_uz.text = uz_title
        p_uz.font.name = FONT_FAMILY
        p_uz.font.size = Pt(14)
        p_uz.font.italic = True
        p_uz.font.color.rgb = ACCENT_CYAN

        # Divider
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(1.68), Inches(11.533), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = DIVIDER
        line.line.color.rgb = DIVIDER

    # =============================================================
    # SLIDE 1: HERO
    # =============================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide1)

    t_box = slide1.shapes.add_textbox(Inches(1.0), Inches(1.3), Inches(6.5), Inches(5.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

    p_tag = tf1.paragraphs[0]
    p_tag.text = "DIGITAL LIFE · UNIVERSITY PRESENTATION"
    p_tag.font.name = FONT_FAMILY
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_BLUE
    p_tag.space_after = Pt(20)

    p1 = tf1.add_paragraph()
    p1.text = "Social Networks\nin My Life"
    p1.font.name = FONT_FAMILY
    p1.font.size = Pt(48)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_EN
    p1.space_after = Pt(8)

    p2 = tf1.add_paragraph()
    p2.text = "Ijtimoiy tarmoqlar mening hayotimda"
    p2.font.name = FONT_FAMILY
    p2.font.size = Pt(24)
    p2.font.italic = True
    p2.font.color.rgb = ACCENT_CYAN
    p2.space_after = Pt(28)

    p3 = tf1.add_paragraph()
    p3.text = "How social media shapes the way we communicate, learn and live"
    p3.font.name = FONT_FAMILY
    p3.font.size = Pt(16)
    p3.font.bold = True
    p3.font.color.rgb = RGBColor(226, 232, 240)
    p3.space_after = Pt(4)

    p4 = tf1.add_paragraph()
    p4.text = "Ijtimoiy tarmoqlar muloqot, ta’lim va turmush tarzimizni qanday shakllantiradi"
    p4.font.name = FONT_FAMILY
    p4.font.size = Pt(13)
    p4.font.italic = True
    p4.font.color.rgb = TEXT_UZ
    p4.space_after = Pt(36)

    p5 = tf1.add_paragraph()
    p5.text = "English · O‘zbekcha   •   01 / 09"
    p5.font.name = FONT_FAMILY
    p5.font.size = Pt(11)
    p5.font.color.rgb = RGBColor(100, 116, 139)

    hero_img = os.path.join(ASSETS_DIR, "hero_phone.png")
    if os.path.exists(hero_img):
        # Set height to 6.7 inches and position closer to text for a connected hero composition
        slide1.shapes.add_picture(hero_img, Inches(7.5), Inches(0.4), height=Inches(6.7))

    # =============================================================
    # SLIDE 2: WHAT ARE SOCIAL NETWORKS?
    # =============================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide2)
    add_slide_header(slide2, 2, "What Are Social Networks?", "Ijtimoiy tarmoqlar nima?")

    left_w = Inches(5.6)
    c_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.8), left_w, Inches(5.3))
    c_box.fill.solid()
    c_box.fill.fore_color.rgb = CARD_BG
    c_box.line.color.rgb = CARD_BORDER
    c_box.line.width = Pt(1)

    tf_c = c_box.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = tf_c.margin_right = Inches(0.35)
    tf_c.margin_top = Inches(0.25)
    tf_c.margin_bottom = Inches(0.15)

    p_d1 = tf_c.paragraphs[0]
    p_d1.text = "Social networks are online platforms where people can communicate, share information and create content. They have become an important part of modern life."
    p_d1.font.name = FONT_FAMILY
    p_d1.font.size = Pt(15)
    p_d1.font.bold = True
    p_d1.font.color.rgb = TEXT_EN
    p_d1.space_after = Pt(5)

    p_d2 = tf_c.add_paragraph()
    p_d2.text = "Ijtimoiy tarmoqlar — odamlar muloqot qilishi, ma’lumot almashishi va kontent yaratishi mumkin bo‘lgan onlayn platformalardir. Ular zamonaviy hayotning muhim qismiga aylandi."
    p_d2.font.name = FONT_FAMILY
    p_d2.font.size = Pt(11.5)
    p_d2.font.italic = True
    p_d2.font.color.rgb = TEXT_UZ
    p_d2.space_after = Pt(16)

    p_c1 = tf_c.add_paragraph()
    p_c1.text = "1. Communication"
    p_c1.font.name = FONT_FAMILY
    p_c1.font.size = Pt(13.5)
    p_c1.font.bold = True
    p_c1.font.color.rgb = ACCENT_BLUE
    p_c1.space_after = Pt(1)

    p_c1_u = tf_c.add_paragraph()
    p_c1_u.text = "   Muloqot — connecting instantly with friends and family"
    p_c1_u.font.name = FONT_FAMILY
    p_c1_u.font.size = Pt(10.5)
    p_c1_u.font.italic = True
    p_c1_u.font.color.rgb = TEXT_UZ
    p_c1_u.space_after = Pt(8)

    p_c2 = tf_c.add_paragraph()
    p_c2.text = "2. Information"
    p_c2.font.name = FONT_FAMILY
    p_c2.font.size = Pt(13.5)
    p_c2.font.bold = True
    p_c2.font.color.rgb = ACCENT_CYAN
    p_c2.space_after = Pt(1)

    p_c2_u = tf_c.add_paragraph()
    p_c2_u.text = "   Ma’lumot — accessing news, updates and global awareness"
    p_c2_u.font.name = FONT_FAMILY
    p_c2_u.font.size = Pt(10.5)
    p_c2_u.font.italic = True
    p_c2_u.font.color.rgb = TEXT_UZ
    p_c2_u.space_after = Pt(8)

    p_c3 = tf_c.add_paragraph()
    p_c3.text = "3. Content sharing"
    p_c3.font.name = FONT_FAMILY
    p_c3.font.size = Pt(13.5)
    p_c3.font.bold = True
    p_c3.font.color.rgb = RGBColor(168, 85, 247)
    p_c3.space_after = Pt(1)

    p_c3_u = tf_c.add_paragraph()
    p_c3_u.text = "   Kontent almashish — publishing creative photos, videos and ideas"
    p_c3_u.font.name = FONT_FAMILY
    p_c3_u.font.size = Pt(10.5)
    p_c3_u.font.italic = True
    p_c3_u.font.color.rgb = TEXT_UZ

    net_img = os.path.join(ASSETS_DIR, "network_visual.png")
    if os.path.exists(net_img):
        slide2.shapes.add_picture(net_img, Inches(6.8), Inches(1.8), width=Inches(5.6))

    # =============================================================
    # SLIDE 3: SOCIAL NETWORKS I USE
    # =============================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide3)
    add_slide_header(slide3, 3, "Social Networks I Use", "Men foydalanadigan ijtimoiy tarmoqlar")

    left_narrative = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.8), Inches(5.3), Inches(5.3))
    left_narrative.fill.solid()
    left_narrative.fill.fore_color.rgb = CARD_BG
    left_narrative.line.color.rgb = CARD_BORDER
    left_narrative.line.width = Pt(1)

    tf_ln = left_narrative.text_frame
    tf_ln.word_wrap = True
    tf_ln.margin_left = tf_ln.margin_right = Inches(0.35)
    tf_ln.margin_top = Inches(0.25)
    tf_ln.margin_bottom = Inches(0.15)

    p_lnt = tf_ln.paragraphs[0]
    p_lnt.text = "POPULAR STUDENT PLATFORMS"
    p_lnt.font.name = FONT_FAMILY
    p_lnt.font.size = Pt(9.5)
    p_lnt.font.bold = True
    p_lnt.font.color.rgb = ACCENT_BLUE
    p_lnt.space_after = Pt(8)

    p_lne = tf_ln.add_paragraph()
    p_lne.text = "Students commonly rely on several established platforms for their daily routines, studies, and creative expression."
    p_lne.font.name = FONT_FAMILY
    p_lne.font.size = Pt(14.5)
    p_lne.font.bold = True
    p_lne.font.color.rgb = TEXT_EN
    p_lne.space_after = Pt(5)

    p_lnu = tf_ln.add_paragraph()
    p_lnu.text = "Talabalar o‘zlarining kundalik faoliyati, ta’limi va ijodiy intilishlarida bir nechta asosiy platformalardan foydalanadilar."
    p_lnu.font.name = FONT_FAMILY
    p_lnu.font.size = Pt(11)
    p_lnu.font.italic = True
    p_lnu.font.color.rgb = TEXT_UZ
    p_lnu.space_after = Pt(14)

    platforms_text = [
        ("Telegram", "communication and useful information", "muloqot va foydali ma’lumotlar uchun", ACCENT_BLUE),
        ("Instagram", "photos, videos and creative content", "fotosuratlar, videolar va ijodiy kontent uchun", RGBColor(225, 48, 108)),
        ("YouTube", "education and entertainment", "ta’lim va ko‘ngilochar kontent uchun", RGBColor(239, 68, 68))
    ]
    for name, en_desc, uz_desc, col in platforms_text:
        p_pn = tf_ln.add_paragraph()
        p_pn.text = f"• {name} — {en_desc}"
        p_pn.font.name = FONT_FAMILY
        p_pn.font.size = Pt(12)
        p_pn.font.bold = True
        p_pn.font.color.rgb = TEXT_EN
        p_pn.space_after = Pt(1)

        p_pnu = tf_ln.add_paragraph()
        p_pnu.text = f"   {uz_desc}"
        p_pnu.font.name = FONT_FAMILY
        p_pnu.font.size = Pt(10)
        p_pnu.font.italic = True
        p_pnu.font.color.rgb = TEXT_UZ
        p_pnu.space_after = Pt(6)

    plat_img = os.path.join(ASSETS_DIR, "platforms_visual.png")
    if os.path.exists(plat_img):
        slide3.shapes.add_picture(plat_img, Inches(6.5), Inches(1.8), width=Inches(5.9))

    # =============================================================
    # SLIDE 4: WHY DO I USE SOCIAL NETWORKS?
    # =============================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide4)
    add_slide_header(slide4, 4, "Why Do I Use Social Networks?", "Nega ijtimoiy tarmoqlardan foydalanaman?")

    s_banner = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.85), Inches(11.533), Inches(1.4))
    s_banner.fill.solid()
    s_banner.fill.fore_color.rgb = CARD_BG
    s_banner.line.color.rgb = CARD_BORDER
    s_banner.line.width = Pt(1)

    tf_sb = s_banner.text_frame
    tf_sb.word_wrap = True
    tf_sb.margin_left = tf_sb.margin_right = Inches(0.4)
    tf_sb.margin_top = Inches(0.2)

    p_sbe = tf_sb.paragraphs[0]
    p_sbe.text = "I use social networks for several reasons. They help me communicate with other people, find useful information, learn new things and relax in my free time."
    p_sbe.font.name = FONT_FAMILY
    p_sbe.font.size = Pt(15.5)
    p_sbe.font.bold = True
    p_sbe.font.color.rgb = TEXT_EN
    p_sbe.space_after = Pt(4)

    p_sbu = tf_sb.add_paragraph()
    p_sbu.text = "Men ijtimoiy tarmoqlardan bir nechta sabablar tufayli foydalanaman. Ular menga boshqa odamlar bilan muloqot qilish, foydali ma’lumot topish, yangi narsalarni o‘rganish va bo‘sh vaqtimda dam olishga yordam beradi."
    p_sbu.font.name = FONT_FAMILY
    p_sbu.font.size = Pt(12)
    p_sbu.font.italic = True
    p_sbu.font.color.rgb = TEXT_UZ

    cats = [
        ("Communication", "Muloqot", "Staying connected with friends, family, and classmates seamlessly.", "Do‘stlar, oila va kursdoshlar bilan oson va yaqin aloqada bo‘lish.", ACCENT_BLUE),
        ("Information", "Ma’lumot", "Discovering news, campus updates, and global affairs.", "Yangiliklar, universitet xabarlari va jahon voqealarini bilib borish.", ACCENT_CYAN),
        ("Learning", "O‘rganish", "Acquiring practical skills, foreign languages, and study advice.", "Amaliy ko‘nikmalar, chet tillari va o‘quv maslahatlarini o‘zlashtirish.", RGBColor(168, 85, 247)),
        ("Entertainment", "Ko‘ngilochar", "Listening to favorite music, watching videos, and resting peacefully.", "Sevimli musiqa tinglash, videolar ko‘rish va maroqli dam olish.", ACCENT_GREEN)
    ]

    cw = Inches(2.72)
    cgap = Inches(0.21)
    for i, (en_t, uz_t, en_d, uz_d, col) in enumerate(cats):
        cx = Inches(0.9) + i * (cw + cgap)
        c_shape = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, Inches(3.45), cw, Inches(3.55))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = CARD_BG
        c_shape.line.color.rgb = col
        c_shape.line.width = Pt(1.5)

        tf_c = c_shape.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_right = Inches(0.25)
        tf_c.margin_top = Inches(0.3)

        p_cn = tf_c.paragraphs[0]
        p_cn.text = f"0{i+1}"
        p_cn.font.name = FONT_FAMILY
        p_cn.font.size = Pt(18)
        p_cn.font.bold = True
        p_cn.font.color.rgb = col
        p_cn.space_after = Pt(12)

        p_ce = tf_c.add_paragraph()
        p_ce.text = en_t
        p_ce.font.name = FONT_FAMILY
        p_ce.font.size = Pt(18)
        p_ce.font.bold = True
        p_ce.font.color.rgb = TEXT_EN
        p_ce.space_after = Pt(2)

        p_cu = tf_c.add_paragraph()
        p_cu.text = uz_t
        p_cu.font.name = FONT_FAMILY
        p_cu.font.size = Pt(13)
        p_cu.font.italic = True
        p_cu.font.color.rgb = col
        p_cu.space_after = Pt(14)

        p_de = tf_c.add_paragraph()
        p_de.text = en_d
        p_de.font.name = FONT_FAMILY
        p_de.font.size = Pt(11)
        p_de.font.color.rgb = RGBColor(226, 232, 240)
        p_de.space_after = Pt(6)

        p_du = tf_c.add_paragraph()
        p_du.text = uz_d
        p_du.font.name = FONT_FAMILY
        p_du.font.size = Pt(10)
        p_du.font.italic = True
        p_du.font.color.rgb = TEXT_UZ

    # =============================================================
    # SLIDE 5: COMMUNICATION
    # =============================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide5)
    add_slide_header(slide5, 5, "Communication with People", "Odamlar bilan muloqot")

    comm_box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.8), Inches(5.3), Inches(5.3))
    comm_box.fill.solid()
    comm_box.fill.fore_color.rgb = CARD_BG
    comm_box.line.color.rgb = CARD_BORDER
    comm_box.line.width = Pt(1)

    tf_cb = comm_box.text_frame
    tf_cb.word_wrap = True
    tf_cb.margin_left = tf_cb.margin_right = Inches(0.35)
    tf_cb.margin_top = Inches(0.25)
    tf_cb.margin_bottom = Inches(0.15)

    p_cbt = tf_cb.paragraphs[0]
    p_cbt.text = "CORE STORY • SPEED & ACCESSIBILITY"
    p_cbt.font.name = FONT_FAMILY
    p_cbt.font.size = Pt(9.5)
    p_cbt.font.bold = True
    p_cbt.font.color.rgb = ACCENT_BLUE
    p_cbt.space_after = Pt(8)

    p_cbe = tf_cb.add_paragraph()
    p_cbe.text = "Social networks make communication faster and easier. I can send messages, share photos and talk to friends and family even when they are far away."
    p_cbe.font.name = FONT_FAMILY
    p_cbe.font.size = Pt(15)
    p_cbe.font.bold = True
    p_cbe.font.color.rgb = TEXT_EN
    p_cbe.space_after = Pt(5)

    p_cbu = tf_cb.add_paragraph()
    p_cbu.text = "Ijtimoiy tarmoqlar muloqotni tezroq va osonroq qiladi. Men do‘stlarim va oilam uzoqda bo‘lsa ham, xabar yuborishim, fotosurat ulashishim va ular bilan suhbatlashishim mumkin."
    p_cbu.font.name = FONT_FAMILY
    p_cbu.font.size = Pt(11)
    p_cbu.font.italic = True
    p_cbu.font.color.rgb = TEXT_UZ
    p_cbu.space_after = Pt(14)

    points_comm = [
        ("Instant Connection Across Borders", "Chegarasiz tezkor aloqa"),
        ("Sharing Daily Moments & Photos", "Kundalik muhim lahzalarni ulashish"),
        ("Face-to-Face Video Conversations", "Yuzma-yuz bevosita video muloqot")
    ]
    for en_p, uz_p in points_comm:
        p_pe = tf_cb.add_paragraph()
        p_pe.text = f"✔  {en_p}"
        p_pe.font.name = FONT_FAMILY
        p_pe.font.size = Pt(12)
        p_pe.font.bold = True
        p_pe.font.color.rgb = TEXT_EN
        p_pe.space_after = Pt(1)

        p_pu = tf_cb.add_paragraph()
        p_pu.text = f"     {uz_p}"
        p_pu.font.name = FONT_FAMILY
        p_pu.font.size = Pt(10)
        p_pu.font.italic = True
        p_pu.font.color.rgb = TEXT_UZ
        p_pu.space_after = Pt(5)

    chat_img = os.path.join(ASSETS_DIR, "chat_visual.png")
    if os.path.exists(chat_img):
        slide5.shapes.add_picture(chat_img, Inches(6.5), Inches(1.8), width=Inches(5.9))

    # =============================================================
    # SLIDE 6: LEARNING AND EDUCATION
    # =============================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide6)
    add_slide_header(slide6, 6, "Learning and Education", "O‘rganish va ta’lim")

    edu_box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.8), Inches(5.3), Inches(5.3))
    edu_box.fill.solid()
    edu_box.fill.fore_color.rgb = CARD_BG
    edu_box.line.color.rgb = CARD_BORDER
    edu_box.line.width = Pt(1)

    tf_eb = edu_box.text_frame
    tf_eb.word_wrap = True
    tf_eb.margin_left = tf_eb.margin_right = Inches(0.35)
    tf_eb.margin_top = Inches(0.25)
    tf_eb.margin_bottom = Inches(0.15)

    p_ebt = tf_eb.paragraphs[0]
    p_ebt.text = "ACADEMIC & DIGITAL LEARNING"
    p_ebt.font.name = FONT_FAMILY
    p_ebt.font.size = Pt(9.5)
    p_ebt.font.bold = True
    p_ebt.font.color.rgb = ACCENT_BLUE
    p_ebt.space_after = Pt(8)

    p_ebe = tf_eb.add_paragraph()
    p_ebe.text = "Social networks can also be useful for education. I can watch educational videos, follow useful channels and find new information about different topics."
    p_ebe.font.name = FONT_FAMILY
    p_ebe.font.size = Pt(14.5)
    p_ebe.font.bold = True
    p_ebe.font.color.rgb = TEXT_EN
    p_ebe.space_after = Pt(5)

    p_ebu = tf_eb.add_paragraph()
    p_ebu.text = "Ijtimoiy tarmoqlar ta’lim uchun ham foydali bo‘lishi mumkin. Men ta’limiy videolar ko‘rishim, foydali kanallarni kuzatishim va turli mavzular haqida yangi ma’lumot topishim mumkin."
    p_ebu.font.name = FONT_FAMILY
    p_ebu.font.size = Pt(11)
    p_ebu.font.italic = True
    p_ebu.font.color.rgb = TEXT_UZ
    p_ebu.space_after = Pt(12)

    edu_elements = [
        ("Video lessons", "video darslar"),
        ("Online courses", "onlayn kurslar"),
        ("Educational channels", "ta’limiy kanallar"),
        ("Useful information", "foydali ma’lumotlar")
    ]
    for en_el, uz_el in edu_elements:
        p_ee = tf_eb.add_paragraph()
        p_ee.text = f"•  {en_el}"
        p_ee.font.name = FONT_FAMILY
        p_ee.font.size = Pt(11.5)
        p_ee.font.bold = True
        p_ee.font.color.rgb = TEXT_EN
        p_ee.space_after = Pt(1)

        p_eu = tf_eb.add_paragraph()
        p_eu.text = f"    {uz_el}"
        p_eu.font.name = FONT_FAMILY
        p_eu.font.size = Pt(9.5)
        p_eu.font.italic = True
        p_eu.font.color.rgb = TEXT_UZ
        p_eu.space_after = Pt(3)

    edu_img = os.path.join(ASSETS_DIR, "education_player.png")
    if os.path.exists(edu_img):
        slide6.shapes.add_picture(edu_img, Inches(6.5), Inches(1.8), width=Inches(5.9))

    # =============================================================
    # SLIDE 7: ENTERTAINMENT AND FREE TIME
    # =============================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide7)
    add_slide_header(slide7, 7, "Entertainment and Free Time", "Ko‘ngilochar va bo‘sh vaqt")

    ent_box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.8), Inches(5.3), Inches(5.3))
    ent_box.fill.solid()
    ent_box.fill.fore_color.rgb = CARD_BG
    ent_box.line.color.rgb = CARD_BORDER
    ent_box.line.width = Pt(1)

    tf_ent = ent_box.text_frame
    tf_ent.word_wrap = True
    tf_ent.margin_left = tf_ent.margin_right = Inches(0.35)
    tf_ent.margin_top = Inches(0.25)
    tf_ent.margin_bottom = Inches(0.15)

    p_ent_t = tf_ent.paragraphs[0]
    p_ent_t.text = "RECREATION & DIGITAL WELLBEING"
    p_ent_t.font.name = FONT_FAMILY
    p_ent_t.font.size = Pt(9.5)
    p_ent_t.font.bold = True
    p_ent_t.font.color.rgb = ACCENT_BLUE
    p_ent_t.space_after = Pt(8)

    p_ente = tf_ent.add_paragraph()
    p_ente.text = "Social networks are also a way to relax. I can watch videos, listen to music, see interesting content and discover new ideas. However, I try not to spend too much time online."
    p_ente.font.name = FONT_FAMILY
    p_ente.font.size = Pt(14.5)
    p_ente.font.bold = True
    p_ente.font.color.rgb = TEXT_EN
    p_ente.space_after = Pt(5)

    p_entu = tf_ent.add_paragraph()
    p_entu.text = "Ijtimoiy tarmoqlar dam olish usuli hamdir. Men videolar ko‘rishim, musiqa tinglashim, qiziqarli kontent tomosha qilishim va yangi g‘oyalarni kashf qilishim mumkin. Biroq internetda juda ko‘p vaqt sarflamaslikka harakat qilaman."
    p_entu.font.name = FONT_FAMILY
    p_entu.font.size = Pt(11)
    p_entu.font.italic = True
    p_entu.font.color.rgb = TEXT_UZ
    p_entu.space_after = Pt(14)

    b_items = [
        ("Relaxation & Music", "Dam olish va musiqa"),
        ("Inspiration & Ideas", "Ilhom va yangi g‘oyalar"),
        ("Healthy Screen Limits", "Sog‘lom ekran vaqti me’yori")
    ]
    for en_bi, uz_bi in b_items:
        p_bie = tf_ent.add_paragraph()
        p_bie.text = f"✔  {en_bi}"
        p_bie.font.name = FONT_FAMILY
        p_bie.font.size = Pt(12)
        p_bie.font.bold = True
        p_bie.font.color.rgb = TEXT_EN
        p_bie.space_after = Pt(1)

        p_biu = tf_ent.add_paragraph()
        p_biu.text = f"     {uz_bi}"
        p_biu.font.name = FONT_FAMILY
        p_biu.font.size = Pt(10)
        p_biu.font.italic = True
        p_biu.font.color.rgb = TEXT_UZ
        p_biu.space_after = Pt(5)

    wb_img = os.path.join(ASSETS_DIR, "balance_wellbeing.png")
    if os.path.exists(wb_img):
        slide7.shapes.add_picture(wb_img, Inches(6.5), Inches(1.8), width=Inches(5.9))

    # =============================================================
    # SLIDE 8: ADVANTAGES VS DISADVANTAGES
    # =============================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide8)
    add_slide_header(slide8, 8, "Advantages and Disadvantages", "Afzalliklari va kamchiliklari")

    col_w = Inches(5.6)
    
    # Left Column: ADVANTAGES
    adv_box = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.8), col_w, Inches(5.3))
    adv_box.fill.solid()
    adv_box.fill.fore_color.rgb = CARD_BG
    adv_box.line.color.rgb = ACCENT_GREEN
    adv_box.line.width = Pt(1.5)

    tf_a = adv_box.text_frame
    tf_a.word_wrap = True
    tf_a.margin_left = tf_a.margin_right = Inches(0.35)
    tf_a.margin_top = Inches(0.25)
    tf_a.margin_bottom = Inches(0.15)

    p_ah = tf_a.paragraphs[0]
    p_ah.text = "ADVANTAGES"
    p_ah.font.name = FONT_FAMILY
    p_ah.font.size = Pt(18)
    p_ah.font.bold = True
    p_ah.font.color.rgb = ACCENT_GREEN
    p_ah.space_after = Pt(1)

    p_ahu = tf_a.add_paragraph()
    p_ahu.text = "AFZALLIKLARI"
    p_ahu.font.name = FONT_FAMILY
    p_ahu.font.size = Pt(12)
    p_ahu.font.italic = True
    p_ahu.font.color.rgb = RGBColor(167, 243, 208)
    p_ahu.space_after = Pt(14)

    adv_items = [
        ("Easy communication", "Oson muloqot"),
        ("Quick access to information", "Ma’lumotga tezkor kirish"),
        ("Educational opportunities", "Ta’lim imkoniyatlari"),
        ("Entertainment", "Ko‘ngilochar imkoniyatlar")
    ]
    for en_i, uz_i in adv_items:
        p_ai = tf_a.add_paragraph()
        p_ai.text = f"✔  {en_i}"
        p_ai.font.name = FONT_FAMILY
        p_ai.font.size = Pt(13)
        p_ai.font.bold = True
        p_ai.font.color.rgb = TEXT_EN
        p_ai.space_after = Pt(1)

        p_aiu = tf_a.add_paragraph()
        p_aiu.text = f"     {uz_i}"
        p_aiu.font.name = FONT_FAMILY
        p_aiu.font.size = Pt(10.5)
        p_aiu.font.italic = True
        p_aiu.font.color.rgb = TEXT_UZ
        p_aiu.space_after = Pt(8)

    # Right Column: DISADVANTAGES
    dis_box = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), col_w, Inches(5.3))
    dis_box.fill.solid()
    dis_box.fill.fore_color.rgb = CARD_BG
    dis_box.line.color.rgb = ACCENT_ROSE
    dis_box.line.width = Pt(1.5)

    tf_d = dis_box.text_frame
    tf_d.word_wrap = True
    tf_d.margin_left = tf_d.margin_right = Inches(0.35)
    tf_d.margin_top = Inches(0.25)
    tf_d.margin_bottom = Inches(0.15)

    p_dh = tf_d.paragraphs[0]
    p_dh.text = "DISADVANTAGES"
    p_dh.font.name = FONT_FAMILY
    p_dh.font.size = Pt(18)
    p_dh.font.bold = True
    p_dh.font.color.rgb = ACCENT_ROSE
    p_dh.space_after = Pt(1)

    p_dhu = tf_d.add_paragraph()
    p_dhu.text = "KAMCHILIKLARI"
    p_dhu.font.name = FONT_FAMILY
    p_dhu.font.size = Pt(12)
    p_dhu.font.italic = True
    p_dhu.font.color.rgb = RGBColor(254, 205, 211)
    p_dhu.space_after = Pt(14)

    dis_items = [
        ("Too much screen time", "Ekran qarshisida ortiqcha vaqt"),
        ("Distraction", "Diqqatni chalg‘itishi"),
        ("False information", "Noto‘g‘ri ma’lumotlar"),
        ("Privacy problems", "Maxfiylik muammolari")
    ]
    for en_i, uz_i in dis_items:
        p_di = tf_d.add_paragraph()
        p_di.text = f"✖  {en_i}"
        p_di.font.name = FONT_FAMILY
        p_di.font.size = Pt(13)
        p_di.font.bold = True
        p_di.font.color.rgb = TEXT_EN
        p_di.space_after = Pt(1)

        p_diu = tf_d.add_paragraph()
        p_diu.text = f"     {uz_i}"
        p_diu.font.name = FONT_FAMILY
        p_diu.font.size = Pt(10.5)
        p_diu.font.italic = True
        p_diu.font.color.rgb = TEXT_UZ
        p_diu.space_after = Pt(8)

    # =============================================================
    # SLIDE 9: CONCLUSION
    # =============================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide9)

    tag_box9 = slide9.shapes.add_textbox(Inches(0.9), Inches(0.45), Inches(8.0), Inches(0.3))
    tf_t9 = tag_box9.text_frame
    p_t9 = tf_t9.paragraphs[0]
    p_t9.text = "FINAL THOUGHTS • SLIDE 09 OF 09"
    p_t9.font.name = FONT_FAMILY
    p_t9.font.size = Pt(10)
    p_t9.font.bold = True
    p_t9.font.color.rgb = ACCENT_BLUE

    stmt_box = slide9.shapes.add_textbox(Inches(0.9), Inches(0.8), Inches(11.533), Inches(1.8))
    tf_st = stmt_box.text_frame
    tf_st.word_wrap = True
    tf_st.margin_left = tf_st.margin_top = 0

    p_ste = tf_st.paragraphs[0]
    p_ste.text = "“Social networks are useful when we use them wisely.”"
    p_ste.font.name = FONT_FAMILY
    p_ste.font.size = Pt(32)
    p_ste.font.bold = True
    p_ste.font.color.rgb = TEXT_EN
    p_ste.space_after = Pt(4)

    p_stu = tf_st.add_paragraph()
    p_stu.text = "“Ijtimoiy tarmoqlardan oqilona foydalansak, ular foydali vositaga aylanishi mumkin.”"
    p_stu.font.name = FONT_FAMILY
    p_stu.font.size = Pt(19)
    p_stu.font.italic = True
    p_stu.font.color.rgb = ACCENT_CYAN

    synth_card = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(2.7), Inches(11.533), Inches(2.3))
    synth_card.fill.solid()
    synth_card.fill.fore_color.rgb = CARD_BG
    synth_card.line.color.rgb = CARD_BORDER
    synth_card.line.width = Pt(1)

    tf_sc = synth_card.text_frame
    tf_sc.word_wrap = True
    tf_sc.margin_left = tf_sc.margin_right = Inches(0.5)
    tf_sc.margin_top = Inches(0.3)

    p_sce = tf_sc.paragraphs[0]
    p_sce.text = "Social networks are an important part of modern life. They help me communicate, learn and relax. However, I believe it is important to use them responsibly and keep a healthy balance between online and real life."
    p_sce.font.name = FONT_FAMILY
    p_sce.font.size = Pt(16.5)
    p_sce.font.bold = True
    p_sce.font.color.rgb = TEXT_EN
    p_sce.space_after = Pt(10)

    p_scu = tf_sc.add_paragraph()
    p_scu.text = "Ijtimoiy tarmoqlar zamonaviy hayotning muhim qismidir. Ular menga muloqot qilish, o‘rganish va dam olishga yordam beradi. Biroq ulardan mas’uliyat bilan foydalanish va onlayn hamda real hayot o‘rtasida sog‘lom muvozanatni saqlash muhim deb hisoblayman."
    p_scu.font.name = FONT_FAMILY
    p_scu.font.size = Pt(13)
    p_scu.font.italic = True
    p_scu.font.color.rgb = TEXT_UZ

    ty_box = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(5.2), Inches(11.533), Inches(1.7))
    ty_box.fill.solid()
    ty_box.fill.fore_color.rgb = RGBColor(15, 23, 42)
    ty_box.line.color.rgb = ACCENT_BLUE
    ty_box.line.width = Pt(1.5)

    tf_ty = ty_box.text_frame
    tf_ty.word_wrap = True
    tf_ty.vertical_anchor = MSO_ANCHOR.MIDDLE

    p_tye = tf_ty.paragraphs[0]
    p_tye.text = "Thank you for your attention!"
    p_tye.font.name = FONT_FAMILY
    p_tye.font.size = Pt(28)
    p_tye.font.bold = True
    p_tye.font.color.rgb = TEXT_EN
    p_tye.alignment = PP_ALIGN.CENTER
    p_tye.space_after = Pt(4)

    p_tyu = tf_ty.add_paragraph()
    p_tyu.text = "E’tiboringiz uchun rahmat!"
    p_tyu.font.name = FONT_FAMILY
    p_tyu.font.size = Pt(19)
    p_tyu.font.italic = True
    p_tyu.font.color.rgb = ACCENT_CYAN
    p_tyu.alignment = PP_ALIGN.CENTER

    out_file = os.path.join(WORKSPACE_DIR, "Social_Networks_in_My_Life.pptx")
    prs.save(out_file)
    print(f"Unified Dark Navy PPTX saved to: {out_file}")

if __name__ == "__main__":
    build_presentation()
